# -*- coding: utf-8 -*-
"""Informe semanal en Word (formato institucional: Verdana, banners azules numerados, pie con numeración).

Uso:  python scripts/informe_semanal.py            genera informes/AAAA-MM-DD_informe_semanal.docx
      python scripts/informe_semanal.py --dias 14  cambia el período cubierto

Si existe config/logo.png, se incluye en el encabezado.
"""
import argparse
import os
import sys
from collections import Counter
from datetime import datetime, timedelta

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from actualizar import CONFIG, DATA, RAIZ, ahora, leer_csv, leer_json  # noqa: E402
from alertas import RELEVANTES, candidatos_pendientes, urgencias_por_vencer  # noqa: E402

AZUL = "1F3864"
FUENTE = "Verdana"


def fmt(f):
    try:
        return datetime.strptime(f[:10], "%Y-%m-%d").strftime("%d-%m-%Y")
    except (ValueError, TypeError):
        return f or ""


def sombrear(elemento, color):
    """Sombreado de párrafo o celda (tipo 'clear', no 'solid')."""
    pr = elemento._p.get_or_add_pPr() if hasattr(elemento, "_p") else elemento._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pr.append(shd)


def campo_pagina(parrafo):
    for tipo, texto in (("begin", None), (None, "PAGE"), ("end", None)):
        run = parrafo.add_run()
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        run._r.append(el)


def banner(doc, numero, texto):
    p = doc.add_paragraph()
    sombrear(p, AZUL)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f" {numero}. {texto.upper()}")
    r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.size = Pt(10)


def tabla(doc, columnas, filas, anchos):
    if not filas:
        p = doc.add_paragraph("Sin registros en el período.")
        p.runs[0].italic = True
        return
    t = doc.add_table(rows=1, cols=len(columnas))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (_, titulo) in enumerate(columnas):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(titulo)
        r.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(c, AZUL)
    for f in filas:
        celdas = t.add_row().cells
        for i, (k, _) in enumerate(columnas):
            celdas[i].text = ""
            r = celdas[i].paragraphs[0].add_run(str(f.get(k, "")))
            r.font.size = Pt(8)
    t.autofit = False
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for i, w in enumerate(anchos):
        t.columns[i].width = Cm(w)
    for fila in t.rows:
        for i, w in enumerate(anchos):
            fila.cells[i].width = Cm(w)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dias", type=int, default=7)
    a = ap.parse_args()
    hoy = ahora()
    desde = hoy - timedelta(days=a.dias)
    proyectos = leer_json(os.path.join(DATA, "proyectos.json"), [])
    tit = {p["boletin"]: p.get("nombre_corto") or p["titulo"] for p in proyectos}
    cambios = [c for c in leer_csv(os.path.join(DATA, "cambios.csv")) if c["fecha_deteccion"] >= desde.strftime("%Y-%m-%d %H:%M")]
    meta = leer_json(os.path.join(DATA, "meta.json"), {})

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    estilo = doc.styles["Normal"]
    estilo.font.name = FUENTE
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    estilo.font.size = Pt(9.5)

    # encabezado y pie
    enc = sec.header.paragraphs[0]
    logo = os.path.join(CONFIG, "logo.png")
    if os.path.exists(logo):
        enc.add_run().add_picture(logo, height=Cm(1.4))
        enc.add_run("   ")
    r = enc.add_run("Subsecretaría de Prevención del Delito")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    pie = sec.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = pie.add_run("Página ")
    rr.font.size = Pt(8)
    campo_pagina(pie)

    # título
    p = doc.add_paragraph()
    r = p.add_run("INFORME SEMANAL DE SEGUIMIENTO LEGISLATIVO")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    p = doc.add_paragraph()
    r = p.add_run(f"Seguridad pública y prevención del delito · Período {fmt(desde.strftime('%Y-%m-%d'))} al {fmt(hoy.strftime('%Y-%m-%d'))}")
    r.font.size = Pt(9)

    # 1. resumen
    banner(doc, 1, "Resumen")
    sit = Counter(p["situacion"] for p in proyectos)
    urg = Counter(p["urgencia_vigente"] for p in proyectos if p["situacion"] == "En tramitación")
    rel = [c for c in cambios if c["tipo"] in RELEVANTES]
    lineas = [
        f"Proyectos seguidos: {len(proyectos)}. En tramitación: {sit.get('En tramitación', 0)}; despachados al Ejecutivo: "
        f"{sit.get('Despachado al Ejecutivo', 0)}; publicados como ley: {sit.get('Publicado como ley', 0)}.",
        f"Urgencias vigentes en proyectos en tramitación: discusión inmediata {urg.get('Discusión inmediata', 0)}, "
        f"suma {urg.get('Suma', 0)}, simple {urg.get('Simple', 0)}.",
        f"Cambios relevantes en el período: {len(rel)} (además de {len(cambios) - len(rel)} trámites menores).",
    ]
    for ln in lineas:
        doc.add_paragraph(ln, style="List Bullet")

    # 2. cambios relevantes
    banner(doc, 2, "Cambios relevantes del período")
    filas = [{"fecha": fmt(c["fecha_deteccion"]), "boletin": c["boletin"], "proyecto": tit.get(c["boletin"], ""),
              "cambio": c["tipo"], "detalle": (f"{c['antes']} → {c['despues']}" if c["antes"] else c["despues"])}
             for c in rel if c["tipo"] not in ("Nuevo informe de comisión", "Nuevo boletín de indicaciones (Senado)",
                                               "Nueva votación en Sala del Senado", "Nueva votación en Sala de la Cámara")]
    tabla(doc, [("fecha", "Detectado"), ("boletin", "Boletín"), ("proyecto", "Proyecto"), ("cambio", "Cambio"), ("detalle", "Detalle")],
          filas, [2, 2, 5.5, 2.6, 5])

    # 3. urgencias
    banner(doc, 3, "Urgencias que vencen en los próximos 7 días")
    tabla(doc, [("boletin", "Boletín"), ("proyecto", "Proyecto"), ("urgencia", "Urgencia"), ("vence", "Vence (estimado)")],
          [{**u, "vence": fmt(u["vence"]), "proyecto": u["proyecto"]} for u in urgencias_por_vencer()], [2, 9.5, 3, 2.6])

    # 4. informes, indicaciones y votaciones
    banner(doc, 4, "Informes, indicaciones y votaciones nuevas")
    filas = [{"fecha": fmt(c["fecha_deteccion"]), "boletin": c["boletin"], "proyecto": tit.get(c["boletin"], ""),
              "tipo": c["tipo"].replace("Nuevo ", "").replace("Nueva ", "").capitalize(), "detalle": c["despues"][:150]}
             for c in rel if c["tipo"] in ("Nuevo informe de comisión", "Nuevo boletín de indicaciones (Senado)",
                                           "Nueva votación en Sala del Senado", "Nueva votación en Sala de la Cámara")]
    tabla(doc, [("fecha", "Detectado"), ("boletin", "Boletín"), ("proyecto", "Proyecto"), ("tipo", "Tipo"), ("detalle", "Detalle")],
          filas, [2, 2, 5, 3, 5.1])

    # 5. candidatos
    banner(doc, 5, "Proyectos nuevos pendientes de revisión")
    tabla(doc, [("boletin", "Boletín"), ("titulo", "Proyecto"), ("motivo", "Motivo de la detección")],
          [{**c, "titulo": c["titulo"]} for c in candidatos_pendientes()], [2, 9, 6.1])

    # 6. estado por materia
    banner(doc, 6, "Proyectos seguidos por materia")
    por_mat = Counter((p["clasificacion"], p["materia"]) for p in proyectos if p["situacion"] == "En tramitación")
    tabla(doc, [("clasificacion", "Clasificación"), ("materia", "Materia"), ("n", "En tramitación")],
          [{"clasificacion": k[0], "materia": k[1], "n": v} for k, v in sorted(por_mat.items(), key=lambda kv: (kv[0][0] != "Núcleo", -kv[1]))],
          [3, 11, 3.1])

    # fuentes (nota al pie del documento, sin recuadros)
    doc.add_paragraph()
    fuentes = "; ".join(f"{k}: {v.get('estado', 's/i')} (última consulta exitosa {fmt(v.get('ultima_consulta_exitosa', ''))})"
                        for k, v in meta.get("fuentes", {}).items())
    p = doc.add_paragraph(f"Fuente: registro oficial de tramitación del Senado y datos abiertos de la Cámara de Diputadas y Diputados. "
                          f"Estado de las fuentes: {fuentes}. Plazos de urgencia estimados desde su ingreso.")
    p.runs[0].font.size = Pt(7.5)
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    carpeta = os.path.join(RAIZ, "informes")
    os.makedirs(carpeta, exist_ok=True)
    ruta = os.path.join(carpeta, f"{hoy.strftime('%Y-%m-%d')}_informe_semanal.docx")
    doc.save(ruta)
    print(ruta)


if __name__ == "__main__":
    main()
