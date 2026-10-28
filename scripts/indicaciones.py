# -*- coding: utf-8 -*-
"""Lectura de los boletines de indicaciones del Senado.

Cada indicación numerada empieza con "N.- De/Del <autor>, para <acción> ...".
Una misma frase puede agrupar varias: "4.- De la Honorable Senadora señora Sánchez y
5.- de la Honorable Senadora señora Pascual, para suprimirlo."
"""
import io
import os
import re
import shutil
import subprocess
import tempfile

MARCA = re.compile(r"(?<![\d\w])(\d{1,4}(?:\s*(?:bis|ter|quáter|quater))?[A-Z]{0,2}(?:\s*[a-z]\))?)\s*\.\s*-\s*(?=[Dd]el?\s|[Dd]e\s)")
ENCABEZADO = re.compile(r"^\s*(ART[ÍI]CULO|Art[íi]culo|Disposici[óo]n|DISPOSICI[ÓO]N|T[ÍI]TULO|P[áa]rrafo|N[úu]mero|Inciso|Letra|Literal|Numeral|Encabezamiento|Ep[íi]grafe)\b.{0,90}$")


def texto_de_documento(contenido):
    """Extrae texto de .docx (python-docx) o .doc (LibreOffice si está instalado)."""
    if contenido[:2] == b"PK":
        import docx
        d = docx.Document(io.BytesIO(contenido))
        partes = [p.text for p in d.paragraphs]
        for tabla in d.tables:
            for fila in tabla.rows:
                for celda in fila.cells:
                    partes.append(celda.text)
        return "\n".join(partes)
    if contenido[:4] == b"\xd0\xcf\x11\xe0":                     # Word 97-2003
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise RuntimeError("formato .doc: requiere LibreOffice para convertir")
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "doc.doc")
            with open(ruta, "wb") as f:
                f.write(contenido)
            perfil = "file://" + os.path.join(tmp, "perfil")      # perfil propio: permite conversiones en paralelo
            subprocess.run([soffice, f"-env:UserInstallation={perfil}", "--headless", "--convert-to", "txt:Text (encoded):UTF8", "--outdir", tmp, ruta],
                           check=True, capture_output=True, timeout=180)
            with open(os.path.join(tmp, "doc.txt"), encoding="utf-8-sig", errors="replace") as f:
                return f.read()
    if contenido[:4] == b"%PDF":
        raise RuntimeError("formato PDF no soportado")
    raise RuntimeError("formato no reconocido (posible bloqueo del servidor)")


def tipo_autor(autor):
    a = autor.lower()
    if "presidente de la rep" in a or "presidenta de la rep" in a or "ejecutivo" in a or "excelencia" in a or re.search(r"\bs\.\s?e\.", a):
        return "Ejecutivo"
    if "senador" in a or "diputad" in a:
        return "Parlamentario"
    return "Otro"


def limpiar_autor(autor):
    a = re.sub(r"\s+", " ", autor).strip(" ,.;")
    a = re.sub(r"^(?:[Dd]e las|[Dd]e los|[Dd]e la|[Dd]el|[Dd]e)\s+", "", a)
    if tipo_autor(a) == "Ejecutivo":
        return "Presidente de la República"
    a = re.sub(r"\b(Honorables?|H\.)\s+", "", a)
    a = re.sub(r"\b(Senador(?:a|es|as)?|Diputad[oa]s?)\s+", "", a)
    a = re.sub(r"\bseñor(?:a|es|as)?\s+", "", a)
    return a.strip(" ,")


def leer_indicaciones(texto):
    """Devuelve lista de {numero, autor, tipo_autor, accion, ubicacion}."""
    lineas = texto.replace("\r", "\n").split("\n")
    salida, ubicacion, articulo = [], "", ""
    bloque = []

    def procesar(parrafo, ubic):
        marcas = list(MARCA.finditer(parrafo))
        if not marcas:
            return
        m_para = re.search(r",?\s+para\s+([a-záéíóúñ]+)", parrafo[marcas[-1].end():])
        accion = m_para.group(1) if m_para else ""
        for i, m in enumerate(marcas):
            fin = marcas[i + 1].start() if i + 1 < len(marcas) else len(parrafo)
            seg = parrafo[m.end():fin]
            corte = re.search(r",\s+para\s|\s+para\s+(?:suprimir|sustituir|reemplazar|agregar|intercalar|incorporar|consultar|eliminar|añadir|anteponer)", seg)
            autor = seg[:corte.start()] if corte else re.sub(r"\s+y\s*$", "", seg.strip())
            autor = re.sub(r"\s+y\s*$", "", autor.strip())
            if len(autor) > 250 or not re.match(r"^(?:[Dd]el?|[Dd]e)\s", autor) or tipo_autor(autor) == "Otro":
                continue
            salida.append({"numero": re.sub(r"\s+", " ", m.group(1)), "autor": limpiar_autor(autor), "tipo_autor": tipo_autor(autor),
                           "accion": accion, "ubicacion": ubic})

    for ln in lineas:
        s = ln.strip()
        if not s:
            continue
        if ENCABEZADO.match(s) and not MARCA.search(s):
            if re.match(r"^(ART[ÍI]CULO|Art[íi]culo|Disposici|DISPOSICI|T[ÍI]TULO)", s):
                articulo, ubicacion = s, s
            else:
                ubicacion = f"{articulo} / {s}" if articulo else s
            continue
        if MARCA.match(s):
            procesar(s, ubicacion)
    # depurar duplicados (tablas repetidas) y ordenar
    vistos, unicos = set(), []
    for x in salida:
        if x["numero"] not in vistos:
            vistos.add(x["numero"])
            unicos.append(x)
    return sorted(unicos, key=clave_numero)


def clave_numero(x):
    m = re.match(r"(\d+)(.*)", x["numero"])
    return (int(m.group(1)), m.group(2)) if m else (0, x["numero"])
