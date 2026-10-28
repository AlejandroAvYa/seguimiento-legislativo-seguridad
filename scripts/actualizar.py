# -*- coding: utf-8 -*-
"""Actualiza la base de proyectos de ley de seguridad pública a partir del Senado.

Uso:  python scripts/actualizar.py [--sin-candidatos] [--sin-excel]

Lee  config/catalogo.csv  (boletines a seguir y su clasificación manual)
     data/candidatos.csv  (candidatos ya revisados, con su decisión)
Escribe en data/:
     proyectos.json  proyectos.csv  tramites.csv  urgencias.csv
     cambios.csv  cambios.json  candidatos.csv  meta.json  proyectos.xlsx
y, si hay candidatos nuevos, salida/alerta_candidatos.md (no se versiona).
"""
import argparse
import csv
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import indicaciones as ind_mod  # noqa: E402
import camara  # noqa: E402
import gestion  # noqa: E402
import agenda as agenda_mod  # noqa: E402
import senado  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG = os.path.join(RAIZ, "config")
DATA = os.path.join(RAIZ, "data")
SALIDA = os.path.join(RAIZ, "salida")

INICIO_GOBIERNO = "2026-03-11"          # asunción del actual Gobierno
DIAS_REVISION_CANDIDATOS = 6            # ventana de búsqueda de proyectos nuevos
MAX_ERRORES = 0.3                       # sobre este % de fallas no se escribe nada

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo("America/Santiago")
except Exception:                       # pragma: no cover
    TZ = timezone(timedelta(hours=-3))

PALABRAS_CLAVE = re.compile(
    r"penal|delit|crimen|criminal|seguridad p[uú]blica|seguridad ciudadana|prevenci[oó]n del delito|carabiner|polic[ií]a|"
    r"investigaciones de chile|gendarmer|c[aá]rcel|penitenci|privad[oa]s? de libertad|migra|extranjer|expulsi|frontera|"
    r"terroris|narcotr|estupefac|drogas|lavado de activos|armas|explosiv|flagran|control de identidad|estado[s]? de excepci|"
    r"fuerzas de orden|uso de la fuerza|leg[ií]tima defensa|robo|hurto|recept|homicid|secuestr|extorsi|sicari|encapuch|"
    r"desorden|barricad|vandal|incivil|20\.084|seguridad privada|seguridad municipal|infraestructura cr[ií]tica|"
    r"inteligencia|ADN|crimen organizado|asociaci[oó]n il[ií]cita|trata de personas|tr[aá]fico de|f[uú]tbol|19\.327|"
    r"abigeato|reclutamiento|agente encubierto|prisi[oó]n preventiva|libertad condicional|reincid|ciberdelit|"
    r"protecci[oó]n ciudadana|v[ií]ctimas", re.I)


# ----------------------------------------------------------------- utilidades
def ahora():
    return datetime.now(TZ)


def leer_csv(ruta):
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def escribir_csv(ruta, filas, campos):
    with open(ruta, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        w.writerows(filas)


def escribir_json(ruta, obj):
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1, sort_keys=False)
        f.write("\n")


def leer_json(ruta, defecto):
    if not os.path.exists(ruta):
        return defecto
    with open(ruta, encoding="utf-8") as f:
        return json.load(f)


def contenido(ruta):
    if not os.path.exists(ruta):
        return None
    with open(ruta, "rb") as f:
        return f.read()


def gobierno_de(fecha_iso):
    if fecha_iso >= INICIO_GOBIERNO:
        return "Gob. Kast"
    if fecha_iso >= "2022-03-11":
        return "Gob. Boric"
    if fecha_iso >= "2018-03-11":
        return "Gob. Piñera"
    return "Gob. anterior a 2018"


def accion_gobierno(ficha):
    """Urgencias e indicaciones del actual Gobierno y si es mensaje suyo."""
    urg = [u for u in ficha["urgencias"] if u["ingreso"] >= INICIO_GOBIERNO]
    ind = [x for x in ficha["indicaciones_ejecutivo"] if x >= INICIO_GOBIERNO]
    fi = senado.fecha_iso(ficha["fecha_ingreso"])
    mensaje = ficha["iniciativa"].startswith("Mensaje") and fi >= INICIO_GOBIERNO
    return urg, ind, mensaje


def situacion(f):
    etapa = f["etapa"].lower()
    if f["leynro"]:
        return "Publicado como ley"
    if "archivado" in etapa or "archivado" in f["estado"].lower():
        return "Archivado"
    rechazo = any(re.search(r"Discusi[oó]n general\s*\.\s*Rechazado", t["descripcion"]) for t in f["tramites"])
    if "rechazado" in f["estado"].lower() or (etapa.startswith("tramitación terminada") and rechazo):
        return "Rechazado"
    if "retirado" in f["estado"].lower():
        return "Retirado"
    if etapa.startswith("tramitación terminada") or "aprobacion presidencial" in etapa or "aprobación presidencial" in etapa:
        return "Despachado al Ejecutivo"
    # El Senado a veces no actualiza la etapa tras el despacho: se usa el trámite "Oficio de ley al Ejecutivo"
    if (any(re.search(r"Oficio de ley al Ejecutivo", t["descripcion"]) for t in f["tramites"]) or "oficio de ley al ejecutivo" in f["subetapa"].lower()) \
            and not re.search(r"observaci|veto", etapa):
        return "Despachado al Ejecutivo"
    return "En tramitación"


def ultimo_tramite(f):
    validos = [t for t in f["tramites"] if re.match(r"\d{4}-\d{2}-\d{2}$", t["fecha"]) and t["fecha"] <= ahora().strftime("%Y-%m-%d") and t["descripcion"]]
    if not validos:
        return {"fecha": "", "descripcion": "", "camara": ""}
    t = max(validos, key=lambda t: t["fecha"])   # max devuelve el primero entre empates
    ult = [x for x in validos if x["fecha"] == t["fecha"]][-1]
    return {"fecha": ult["fecha"], "descripcion": ult["descripcion"], "camara": ult["camara"]}


PLAZO_URGENCIA = {"Simple": 30, "Suma": 15, "Discusión inmediata": 6}
HITOS = [
    (r"^Ingreso de proyecto", "Ingreso"),
    (r"Aprobado en general y particular", "Aprobado en general y particular"),
    (r"Discusi[oó]n general\s*\.\s*Aprobado en general(?! y particular)", "Aprobado en general"),
    (r"Discusi[oó]n particular\s*\.\s*Aprobado", "Aprobado en particular"),
    (r"Discusi[oó]n general\s*\.\s*Rechazado", "Rechazado en general"),
    (r"Oficio de ley a C[aá]mara Revisora", "Despachado a la cámara revisora"),
    (r"Oficio (con )?modificaciones a C[aá]mara de Origen", "Despachado con modificaciones a la cámara de origen"),
    (r"Oficio aprobaci[oó]n sin modificaciones a C[aá]mara de Origen", "Aprobado sin modificaciones por la cámara revisora"),
    (r"Oficio rechazo modificaciones", "Rechazo de modificaciones (pasa a Comisión Mixta)"),
    (r"integrantes de Comisi[oó]n Mixta", "Integración de Comisión Mixta"),
    (r"Oficio al Tribunal Constitucional|Oficio N.{0,12}al Tribunal Constitucional", "Enviado al Tribunal Constitucional"),
    (r"insistencia", "Insistencia"),
    (r"Oficio de ley al Ejecutivo", "Despachado al Ejecutivo"),
]


def _dias(desde, hasta):
    try:
        return (datetime.strptime(hasta, "%Y-%m-%d") - datetime.strptime(desde, "%Y-%m-%d")).days
    except ValueError:
        return ""


def estructura_etapa(f):
    """'Segundo trámite constitucional (Senado)' -> ('Segundo trámite constitucional', 'Senado'); instancia actual."""
    m = re.match(r"^(.*?)\s*\(([^)]*)\)\s*$", f["etapa"])
    tramite, camara = (m.group(1), m.group(2)) if m else (f["etapa"], "")
    camara = camara.replace("C.Diputados", "Cámara de Diputados")
    sub = f["subetapa"] if f["subetapa"] not in ("", "/") else ""
    mc = re.search(r"[Cc]omisi[oó]n(?:es)?\s+(?:de\s+|para\s+)?(.+)$", sub)
    if re.search(r"[Cc]omisi[oó]n Mixta", sub):
        instancia = "Comisión Mixta"
    elif mc:
        instancia = "Comisión de " + mc.group(1).strip()
    elif re.search(r"^Discusi[oó]n (general|particular|[uú]nica)|Sala|[Cc]uenta", sub):
        instancia = "Sala"
    else:
        instancia = sub
    return tramite, camara, instancia


def inicio_etapa(f):
    trs = [t for t in f["tramites"] if re.match(r"\d{4}-\d{2}-\d{2}$", t["fecha"])]
    if not trs:
        return ""
    ultima = trs[-1]["etapa"]
    inicio = trs[-1]["fecha"]
    for t in reversed(trs):
        if t["etapa"] != ultima:
            break
        inicio = t["fecha"]
    return inicio


def vencimiento_urgencia(f):
    tipo = f["urgencia_actual"]
    if tipo not in PLAZO_URGENCIA:
        return "", ""
    vig = [u for u in f["urgencias"] if u["tipo"] == tipo and not u["retiro"] and re.match(r"\d{4}-", u["ingreso"])]
    if not vig:
        return "", ""
    desde = vig[-1]["ingreso"]
    vence = (datetime.strptime(desde, "%Y-%m-%d") + timedelta(days=PLAZO_URGENCIA[tipo])).strftime("%Y-%m-%d")
    return desde, vence


def hitos(f):
    vistos, salida = set(), []
    for t in f["tramites"]:
        for patron, nombre in HITOS:
            if re.search(patron, t["descripcion"], re.I):
                clave = (nombre, t["camara"], t["etapa"])
                if clave not in vistos:
                    vistos.add(clave)
                    salida.append({"fecha": t["fecha"], "hito": nombre, "camara": t["camara"], "etapa": t["etapa"]})
                break
    return salida


def items_efectivos(boletines, cache):
    """Indicaciones de cada boletín sin repetir las de un boletín anterior que lo reedita.

    Algunos boletines republican uno anterior completo y agregan indicaciones nuevas. Si un boletín contiene al
    menos el 80 % de las indicaciones de otro anterior del mismo proyecto, solo se cuentan las que no estaban.
    """
    clave = lambda x: (x["numero"], x["autor"], x["accion"], x["ubicacion"])
    previos, salida = [], []
    for bi in sorted(boletines, key=lambda x: (x["fecha"], x["iddocto"])):
        items = cache.get(bi["iddocto"], {}).get("items", [])
        claves = {clave(x) for x in items}
        reed = ""
        for idp, cp in previos:
            if cp and len(cp & claves) >= 0.8 * len(cp):
                reed = idp
                claves_prev = cp
                break
        nuevos = [x for x in items if not reed or clave(x) not in claves_prev]
        salida.append((bi, items, nuevos, reed))
        previos.append((bi["iddocto"], claves))
    return salida


def autores_indicaciones(f, cache):
    conteo = {}
    for _, _, nuevos, _ in items_efectivos(f["boletines_indicaciones"], cache):
        for x in nuevos:
            for a in separar_autores(x["autor"]):
                conteo[(a, x["tipo_autor"])] = conteo.get((a, x["tipo_autor"]), 0) + 1
    return [{"autor": a, "tipo_autor": t, "n": n} for (a, t), n in sorted(conteo.items(), key=lambda kv: (-kv[1], kv[0][0]))]


CAMPOS_GESTION = ["responsable", "posicion_ejecutivo", "ultima_gestion", "proximo_paso"]
CAMPOS_CATALOGO = ["nombre_corto", "clasificacion", "materia", "descripcion", "acot", "nota"] + CAMPOS_GESTION


def articulo_de(ubicacion):
    a = (ubicacion or "").split(" / ")[0].strip()
    return re.sub(r"\s+", " ", a) or "Sin ubicación identificada"


def indicaciones_por_articulo(f, cache):
    """Cuántas indicaciones recibió cada artículo, por tipo de autor y acción (en el orden del boletín)."""
    res = {}
    for _, _, nuevos, _ in items_efectivos(f["boletines_indicaciones"], cache):
        for x in nuevos:
            a = articulo_de(x["ubicacion"])
            d = res.setdefault(a, {"articulo": a, "n": 0, "ejecutivo": 0, "parlamentarias": 0, "acciones": {}})
            d["n"] += 1
            d["ejecutivo"] += x["tipo_autor"] == "Ejecutivo"
            d["parlamentarias"] += x["tipo_autor"] == "Parlamentario"
            acc = x["accion"][:-2] if x["accion"].endswith("lo") else (x["accion"] or "otra")
            d["acciones"][acc] = d["acciones"].get(acc, 0) + 1
    return list(res.values())


def resumen_indicaciones(f, cache):
    """Totales por autor a partir de los boletines de indicaciones del Senado ya leídos."""
    lista = []
    for _, _, nuevos, _ in items_efectivos(f["boletines_indicaciones"], cache):
        lista += nuevos
    eje = sum(x["tipo_autor"] == "Ejecutivo" for x in lista)
    eje_actual = 0
    for bi, _, nuevos, _ in items_efectivos(f["boletines_indicaciones"], cache):
        if bi["fecha"] >= INICIO_GOBIERNO:
            eje_actual += sum(x["tipo_autor"] == "Ejecutivo" for x in nuevos)
    par = sum(x["tipo_autor"] == "Parlamentario" for x in lista)
    oficios_eje = [o for o in f["oficios"] if re.search(r"(Presidente de la Rep[uú]blica|Ejecutivo).{0,40}indicaci", o["descripcion"], re.I)]
    return {
        "n_boletines_indicaciones_senado": len(f["boletines_indicaciones"]),
        "n_indicaciones_senado": len(lista),
        "n_indicaciones_senado_ejecutivo": eje,
        "n_indicaciones_senado_ejecutivo_gobierno_actual": eje_actual,
        "n_indicaciones_senado_parlamentarias": par,
        "n_oficios_indicaciones_ejecutivo": len(oficios_eje),
        "ultima_presentacion_indicaciones": max([b["fecha"] for b in f["boletines_indicaciones"]] + [o["fecha"] for o in oficios_eje] + [""]),
    }


def registro(cat, f, cache=None):
    """Combina la ficha oficial con la clasificación manual del catálogo."""
    urg, ind, mensaje = accion_gobierno(f)
    fi = senado.fecha_iso(f["fecha_ingreso"])
    ini = "Mensaje" if f["iniciativa"].startswith("Mensaje") else "Moción"
    m = re.search(r"iddocto=(\d+)", f["link_mensaje_mocion"])
    b = f["boletin"]
    sub = f["subetapa"] if f["subetapa"] not in ("", "/") else ""
    ult = ultimo_tramite(f)
    hoy = ahora().strftime("%Y-%m-%d")
    tram, cam_act, inst = estructura_etapa(f)
    ini_etapa = inicio_etapa(f)
    urg_desde, urg_vence = vencimiento_urgencia(f)
    return {
        "boletin": b,
        "nombre_corto": cat.get("nombre_corto", "") or f["titulo"][:70],
        "acot_n": "", "acot_iniciativa": "", "acot_grupo": "", "agenda": [], "plazo_indicaciones": [],
        "titulo": f["titulo"],
        "clasificacion": cat.get("clasificacion", ""),
        "materia": cat.get("materia", ""),
        "descripcion": cat.get("descripcion", ""),
        "acot": cat.get("acot", ""),
        "nota": cat.get("nota", ""),
        **{k: cat.get(k, "") for k in CAMPOS_GESTION},
        "iniciativa": ini,
        "autoria": gobierno_de(fi) if ini == "Mensaje" else "; ".join(f["autores"]),
        "fecha_ingreso": fi,
        "camara_origen": f["camara_origen"].replace("C.Diputados", "Cámara de Diputados"),
        "situacion": situacion(f),
        "etapa": f["etapa"],
        "subetapa": sub,
        "tramite_constitucional": tram,
        "camara_actual": cam_act,
        "instancia_actual": inst,
        "fecha_inicio_etapa": ini_etapa,
        "dias_en_etapa": _dias(ini_etapa, hoy) if ini_etapa else "",
        "dias_desde_ingreso": _dias(fi, hoy),
        "urgencia_vigente": f["urgencia_actual"] or "Sin urgencia",
        "urgencia_vigente_desde": urg_desde,
        "urgencia_vence_estimado": urg_vence,
        "urgencia_plazo": ("" if not urg_vence else (f"Vence {urg_vence}" if urg_vence >= hoy else f"Plazo cumplido el {urg_vence} sin renovación registrada")),
        "ley_numero": f["leynro"],
        "ley_fecha_diario_oficial": senado.fecha_iso(f["diariooficial"]),
        "refundidos": f["refundidos"].replace("*matriz*", "(matriz)"),
        "mensaje_gobierno_actual": mensaje,
        "n_urgencias_gobierno_actual": len(urg),
        "primera_urgencia_gobierno_actual": urg[0]["ingreso"] if urg else "",
        "ultima_urgencia": (f"{f['urgencias'][-1]['tipo']} ({f['urgencias'][-1]['ingreso']})" if f["urgencias"] else ""),
        "indicaciones_ejecutivo_gobierno_actual": ", ".join(ind),
        "ultimo_tramite_fecha": ult["fecha"],
        "ultimo_tramite": ult["descripcion"],
        "ultimo_tramite_camara": ult["camara"],
        "n_tramites": len(f["tramites"]),
        **resumen_indicaciones(f, cache or {}),
        "n_informes_comision": len(f["informes"]),
        "ultimo_informe": (f"{f['informes'][-1]['fecha']} – {f['informes'][-1]['tramite']}" if f["informes"] else ""),
        "n_votaciones_sala_senado": len(f["votaciones"]),
        "ultima_votacion_sala_senado": (lambda v: f"{v['fecha']} – {v['tipo']}: {v['resultado'] or 's/r'} ({v['si']}-{v['no']}-{v['abstencion']})")(f["votaciones"][-1]) if f["votaciones"] else "",
        "n_votaciones_sala_camara": 0,
        "ultima_votacion_sala_camara": "",
        "veto": "Sí" if f["observaciones"] else "No",
        "link_senado": f"https://tramitacion.senado.cl/appsenado/templates/tramitacion/index.php?boletin_ini={b}",
        "link_camara": (f"https://www.camara.cl/legislacion/ProyectosDeLey/tramitacion.aspx?prmID={m.group(1)}&prmBOLETIN={b}" if m else ""),
        "link_texto": f["link_mensaje_mocion"],
        "tramites": f["tramites"],
        "urgencias": f["urgencias"],
        "hitos": hitos(f),
        "informes": f["informes"],
        "comparados": f["comparados"],
        "oficios": f["oficios"],
        "boletines_indicaciones": [{**bi, "estado_lectura": (cache or {}).get(bi["iddocto"], {}).get("estado", "pendiente")} for bi in f["boletines_indicaciones"]],
        "votaciones": f["votaciones"],
        "indicaciones_por_autor": autores_indicaciones(f, cache or {}),
        "indicaciones_por_articulo": indicaciones_por_articulo(f, cache or {}),
    }


CAMPOS_PLANOS = ["boletin", "nombre_corto", "acot_n", "acot_iniciativa", "titulo", "clasificacion", "materia", "descripcion", "acot", "nota",
                 "etapa_codigo", "etapa_texto", "prioridad_nivel", "prioridad_indice", "accion_sugerida", "n_alertas",
                 "responsable", "posicion_ejecutivo", "ultima_gestion", "proximo_paso", "iniciativa", "autoria",
                 "fecha_ingreso", "camara_origen", "situacion", "etapa", "subetapa", "tramite_constitucional", "camara_actual",
                 "instancia_actual", "fecha_inicio_etapa", "dias_en_etapa", "dias_desde_ingreso",
                 "urgencia_vigente", "urgencia_vigente_desde", "urgencia_vence_estimado", "urgencia_plazo", "ley_numero",
                 "ley_fecha_diario_oficial", "refundidos", "mensaje_gobierno_actual", "n_urgencias_gobierno_actual",
                 "primera_urgencia_gobierno_actual", "ultima_urgencia", "indicaciones_ejecutivo_gobierno_actual",
                 "ultimo_tramite_fecha", "ultimo_tramite", "ultimo_tramite_camara", "n_tramites",
                 "n_boletines_indicaciones_senado", "n_indicaciones_senado", "n_indicaciones_senado_ejecutivo",
                 "n_indicaciones_senado_parlamentarias", "n_indicaciones_senado_ejecutivo_gobierno_actual", "n_oficios_indicaciones_ejecutivo", "ultima_presentacion_indicaciones",
                 "n_informes_comision", "ultimo_informe", "n_votaciones_sala_senado", "ultima_votacion_sala_senado",
                 "n_votaciones_sala_camara", "ultima_votacion_sala_camara", "veto",
                 "link_senado", "link_camara", "link_texto"]
# versión liviana para la portada del dashboard (la ficha completa se carga desde proyectos.json)
CAMPOS_TABLERO = ["boletin", "nombre_corto", "acot_n", "acot_iniciativa", "acot_grupo", "agenda", "plazo_indicaciones", "titulo", "clasificacion", "materia", "descripcion", "acot", "situacion",
                  "etapa_codigo", "etapa_texto", "avance_estados", "avance_camaras", "instancia_actual", "camara_actual",
                  "camara_origen", "dias_en_etapa", "fecha_ingreso", "iniciativa", "urgencia_vigente", "urgencia_vence_estimado",
                  "urgencia_plazo", "prioridad_indice", "prioridad_nivel", "accion_sugerida", "alertas", "ultimo_tramite_fecha",
                  "ultimo_tramite", "ley_numero", "ley_fecha_diario_oficial", "mensaje_gobierno_actual",
                  "n_urgencias_gobierno_actual", "n_indicaciones_senado", "n_indicaciones_senado_ejecutivo",
                  "n_votaciones_sala_senado", "n_votaciones_sala_camara", "link_senado", "responsable"]
CAMPOS_VIGILADOS = [("situacion", "Situación"), ("etapa", "Etapa"), ("subetapa", "Subetapa"),
                    ("urgencia_vigente", "Urgencia vigente"), ("ley_numero", "N° de ley"), ("refundidos", "Refundidos")]


# ----------------------------------------------------------------- cambios
def detectar_cambios(previos, nuevos, marca, camara_con_base=True):
    """camara_con_base=False en la primera carga exitosa de la Cámara: las votaciones históricas no se avisan como nuevas."""
    cambios = []
    for b, n in nuevos.items():
        p = previos.get(b)
        if p is None:
            cambios.append({"fecha_deteccion": marca, "boletin": b, "tipo": "Proyecto agregado al seguimiento",
                            "antes": "", "despues": n["titulo"]})
            continue
        for campo, etiqueta in CAMPOS_VIGILADOS:
            if (p.get(campo) or "") != (n.get(campo) or ""):
                cambios.append({"fecha_deteccion": marca, "boletin": b, "tipo": etiqueta,
                                "antes": p.get(campo) or "", "despues": n.get(campo) or ""})
        for lista, etiqueta, fmt in [
                ("informes", "Nuevo informe de comisión", lambda x: f"{x['fecha']} – {x['tramite']}"),
                ("boletines_indicaciones", "Nuevo boletín de indicaciones (Senado)", lambda x: f"{x['fecha']} – {x['etapa']}"),
                ("votaciones", "Nueva votación en Sala del Senado", lambda x: f"{x['fecha']} – {x['tipo']}: {x['resultado']} ({x['si']}-{x['no']}-{x['abstencion']})"),
                ("votaciones_camara", "Nueva votación en Sala de la Cámara", lambda x: f"{x['fecha']} – {x['tipo']}: {x['resultado']} ({x['si']}-{x['no']}-{x['abstencion']})")]:
            previas = {fmt(x) for x in p.get(lista, [])}
            if lista == "votaciones_camara" and not camara_con_base:
                continue
            if lista in p:                          # evita falsos avisos al migrar desde la versión anterior
                for x in n.get(lista, []):
                    if fmt(x) not in previas:
                        cambios.append({"fecha_deteccion": marca, "boletin": b, "tipo": etiqueta, "antes": "", "despues": fmt(x)})
        vistos = {(t["fecha"], t["descripcion"]) for t in p.get("tramites", [])}
        for t in n["tramites"]:
            if (t["fecha"], t["descripcion"]) not in vistos and t["descripcion"]:
                cambios.append({"fecha_deteccion": marca, "boletin": b, "tipo": "Nuevo trámite",
                                "antes": "", "despues": f"{t['fecha']} – {t['descripcion']} ({t['camara']})"})
    for b in previos:
        if b not in nuevos:
            cambios.append({"fecha_deteccion": marca, "boletin": b, "tipo": "Proyecto retirado del seguimiento",
                            "antes": previos[b].get("titulo", ""), "despues": ""})
    return cambios


# ----------------------------------------------------------------- candidatos
def buscar_candidatos(seguidos, revisados, marca):
    desde = (ahora() - timedelta(days=DIAS_REVISION_CANDIDATOS)).strftime("%d/%m/%Y")
    nuevos = []
    for f in senado.con_movimiento_desde(desde):
        b = f["boletin"]
        if b in seguidos or b in revisados:
            continue
        urg, ind, mensaje = accion_gobierno(f)
        if not (urg or ind or mensaje):
            continue
        if not PALABRAS_CLAVE.search(f["titulo"]):
            continue
        motivo = []
        if mensaje:
            motivo.append("Mensaje del actual Gobierno")
        if urg:
            motivo.append(f"{len(urg)} urgencia(s) del actual Gobierno (última {urg[-1]['tipo']}, {urg[-1]['ingreso']})")
        if ind:
            motivo.append("Indicaciones del Ejecutivo: " + ", ".join(ind))
        nuevos.append({"boletin": b, "titulo": f["titulo"], "fecha_deteccion": marca[:10],
                       "motivo": "; ".join(motivo), "decision": "pendiente", "observacion": ""})
    return nuevos


# ----------------------------------------------------------------- indicaciones y detalle
CAMPOS_BOL_IND = ["boletin", "iddocto", "fecha", "etapa", "tramite", "estado_lectura", "n_indicaciones", "reedita_boletin",
                  "n_indicaciones_nuevas", "n_ejecutivo", "n_parlamentarias", "link"]
CAMPOS_IND = ["boletin", "iddocto", "fecha_boletin", "etapa", "numero", "autor", "tipo_autor", "accion", "ubicacion", "texto"]


RUTA_CACHE = os.path.join(DATA, "cache_boletines_indicaciones.json")


def cargar_cache_indicaciones():
    """Contenido ya leído de cada boletín de indicaciones (por iddocto), para no descargarlo de nuevo."""
    return leer_json(RUTA_CACHE, {})


def guardar_cache_indicaciones(cache):
    escribir_json(RUTA_CACHE, {k: cache[k] for k in sorted(cache, key=lambda x: int(x) if x.isdigit() else 0)})


def leer_boletin_indicaciones(bi):
    try:
        texto = ind_mod.texto_de_documento(senado.documento(bi["link"]))
        items = ind_mod.leer_indicaciones(texto)
        return {"estado": "leído" if items else "sin indicaciones reconocibles", "items": items}
    except Exception as e:  # noqa: BLE001
        return {"estado": f"no leído: {e}"[:120], "items": []}


def separar_autores(autor):
    if autor == "Presidente de la República":
        return [autor]
    partes = re.split(r",\s*(?:y\s+|e\s+)?|\s+(?:y|e)\s+", autor)
    return [x.strip() for x in partes if x.strip()]


def escribir_detalle(orden, cache):
    bol, ind, por_autor, inf, ofi, vot, votos, hit = [], [], {}, [], [], [], [], []
    for p in orden:
        b = p["boletin"]
        for bi, items, nuevos, reed in items_efectivos(p.get("boletines_indicaciones", []), cache):
            bi = {k: v for k, v in bi.items() if k != "estado_lectura"}
            bol.append({"boletin": b, **bi, "estado_lectura": cache.get(bi["iddocto"], {}).get("estado", "pendiente"),
                        "n_indicaciones": len(items), "reedita_boletin": reed, "n_indicaciones_nuevas": len(nuevos),
                        "n_ejecutivo": sum(x["tipo_autor"] == "Ejecutivo" for x in nuevos),
                        "n_parlamentarias": sum(x["tipo_autor"] == "Parlamentario" for x in nuevos)})
            for x in nuevos:
                ind.append({"boletin": b, "iddocto": bi["iddocto"], "fecha_boletin": bi["fecha"], "etapa": bi["etapa"], **x})
                for a in separar_autores(x["autor"]):
                    k = (b, a, x["tipo_autor"])
                    por_autor[k] = por_autor.get(k, 0) + 1
        inf += [{"boletin": b, **x} for x in p.get("informes", [])]
        ofi += [{"boletin": b, **x} for x in p.get("oficios", [])]
        hit += [{"boletin": b, **x} for x in p.get("hitos", [])]
        for v in p.get("votaciones", []):
            vot.append({"boletin": b, **{k: v[k] for k in v if k != "votos"}})
            votos += [{"boletin": b, "fecha": v["fecha"], "tipo": v["tipo"], **x} for x in v["votos"]]
    carpeta = os.path.join(DATA, "indicaciones")
    os.makedirs(carpeta, exist_ok=True)
    por_bol = {}
    for x in ind:
        por_bol.setdefault(x["boletin"], []).append({k: x[k] for k in CAMPOS_IND if k != "boletin"})
    links = {(x["boletin"], x["iddocto"]): x.get("link", "") for x in bol}
    for b, lista in por_bol.items():
        for x in lista:
            x["link"] = links.get((b, x["iddocto"]), "")
        escribir_json(os.path.join(carpeta, f"{b}.json"), lista)
    for viejo in os.listdir(carpeta):
        if viejo.endswith(".json") and viejo[:-5] not in por_bol:
            os.remove(os.path.join(carpeta, viejo))
    escribir_csv(os.path.join(DATA, "boletines_indicaciones.csv"), bol, CAMPOS_BOL_IND)
    escribir_csv(os.path.join(DATA, "indicaciones.csv"), ind, CAMPOS_IND)
    escribir_csv(os.path.join(DATA, "indicaciones_por_autor.csv"),
                 [{"boletin": k[0], "autor": k[1], "tipo_autor": k[2], "n_indicaciones": v} for k, v in sorted(por_autor.items())],
                 ["boletin", "autor", "tipo_autor", "n_indicaciones"])
    escribir_csv(os.path.join(DATA, "informes.csv"), inf, ["boletin", "fecha", "tramite", "etapa", "link"])
    escribir_csv(os.path.join(DATA, "oficios.csv"), ofi, ["boletin", "fecha", "numero", "tipo", "camara", "tramite", "etapa", "descripcion", "link"])
    escribir_csv(os.path.join(DATA, "hitos.csv"), hit, ["boletin", "fecha", "hito", "camara", "etapa"])
    escribir_csv(os.path.join(DATA, "votaciones.csv"), vot, ["boletin", "fecha", "sesion", "tipo", "etapa", "quorum", "si", "no", "abstencion", "pareo", "resultado", "tema"])
    escribir_csv(os.path.join(DATA, "votos.csv"), votos, ["boletin", "fecha", "tipo", "parlamentario", "seleccion"])
    return {"indicaciones": ind, "votaciones": vot, "informes": inf, "hitos": hit, "boletines_indicaciones": bol}


# ----------------------------------------------------------------- Cámara de Diputados
RUTA_CACHE_CAMARA = os.path.join(DATA, "cache_votaciones_camara.json")
RUTA_AGENDA = os.path.join(DATA, "agenda.json")


def acot_filas():
    return leer_csv(os.path.join(CONFIG, "acot.csv"))


def cargar_acot():
    """Boletín -> N° e iniciativa de la Agenda contra el Crimen Organizado y el Terrorismo (ACOT)."""
    res = {}
    for r in acot_filas():
        for b in [r.get("boletin", "")] + [x for x in r.get("boletines_relacionados", "").split(";") if x]:
            if b.strip():
                res[b.strip()] = {"acot_n": r["n_acot"], "acot_iniciativa": r["iniciativa"], "acot_grupo": r["grupo"]}
    return res


def acot_lista(filas, proyectos):
    """Las iniciativas ACOT con su proyecto, incluidas las anunciadas que aún no tienen boletín."""
    salida = []
    for r in filas:
        b = r.get("boletin", "").strip()
        p = proyectos.get(b, {})
        salida.append({**r, "en_seguimiento": bool(p), "nombre_corto": p.get("nombre_corto", ""),
                       "situacion": p.get("situacion", ""), "etapa_texto": p.get("etapa_texto", "")})
    return salida
PRIMER_ANNO_CAMARA = 2022


def actualizar_camara(boletines):
    """Votaciones de Sala de la Cámara para los boletines seguidos. Guarda lo descargado en caché.

    Primera vez (o si se agregan boletines nuevos al catálogo): revisa desde 2022. Luego, solo el año en curso
    (y el anterior durante enero). Devuelve (cache, n_errores_detalle).
    """
    cache = leer_json(RUTA_CACHE_CAMARA, {"cubiertos": [], "votaciones": {}, "detalles": {}})
    hoy = ahora()
    nuevos_bol = set(boletines) - set(cache["cubiertos"])
    if nuevos_bol:
        annos = list(range(PRIMER_ANNO_CAMARA, hoy.year + 1))
    else:
        annos = [hoy.year] + ([hoy.year - 1] if hoy.month == 1 else [])
    for a in annos:
        for v in camara.votaciones_del_anno(a):
            if set(v["boletines"]) & set(boletines):
                cache["votaciones"][v["id"]] = v
    cache["cubiertos"] = sorted(set(cache["cubiertos"]) | set(boletines))
    errores = 0
    pendientes = [vid for vid in cache["votaciones"] if vid not in cache["detalles"]]

    def det(vid):
        try:
            return vid, camara.detalle_votacion(vid)
        except Exception:  # noqa: BLE001
            return vid, None
    with ThreadPoolExecutor(max_workers=3) as ex:
        for vid, votos in ex.map(det, pendientes):
            if votos is None:
                errores += 1
            else:
                cache["detalles"][vid] = votos
    escribir_json(RUTA_CACHE_CAMARA, {"cubiertos": cache["cubiertos"],
                                      "votaciones": dict(sorted(cache["votaciones"].items())),
                                      "detalles": dict(sorted(cache["detalles"].items()))})
    return cache, errores


def votaciones_camara_de(boletin, cache):
    salida = []
    for vid, v in cache.get("votaciones", {}).items():
        if boletin in v["boletines"]:
            salida.append({**{k: v[k] for k in v if k != "boletines"}, "votos": cache.get("detalles", {}).get(vid, [])})
    return sorted(salida, key=lambda x: (x["fecha"], x.get("hora", ""), x["id"]))


def agregar_camara(registro_p, cache):
    vc = votaciones_camara_de(registro_p["boletin"], cache)
    registro_p["votaciones_camara"] = vc
    registro_p["n_votaciones_sala_camara"] = len(vc)
    registro_p["ultima_votacion_sala_camara"] = (
        f"{vc[-1]['fecha']} – {vc[-1]['tipo']}: {vc[-1]['resultado'] or 's/r'} ({vc[-1]['si']}-{vc[-1]['no']}-{vc[-1]['abstencion']})" if vc else "")


# ----------------------------------------------------------------- actores
def separar_nombre_voto(n):
    return re.sub(r"\s+", " ", n).strip()


def calcular_actores(orden):
    """Perfiles básicos por fuente, solo sobre los proyectos seguidos (sin cruzar identidades entre fuentes)."""
    autores, indic, votos_s, votos_c = {}, {}, {}, {}
    for p in orden:
        if p["iniciativa"] == "Moción":
            for a in [x.strip() for x in p["autoria"].split(";") if x.strip()]:
                autores.setdefault(a, set()).add(p["boletin"])
        for a in p.get("indicaciones_por_autor", []):
            if a["tipo_autor"] == "Parlamentario" or a["autor"] == "Presidente de la República":
                d = indic.setdefault(a["autor"], {"n": 0, "proyectos": set()})
                d["n"] += a["n"]
                d["proyectos"].add(p["boletin"])
        for lista, destino, campo in ((p.get("votaciones", []), votos_s, "parlamentario"), (p.get("votaciones_camara", []), votos_c, "diputado")):
            for v in lista:
                for x in v.get("votos", []):
                    nombre = separar_nombre_voto(x.get(campo, ""))
                    if not nombre:
                        continue
                    d = destino.setdefault(nombre, {"si": 0, "no": 0, "abstencion": 0, "otros": 0, "proyectos": set()})
                    sel = (x.get("seleccion") or "").lower()
                    k = "si" if sel in ("si", "sí", "afirmativo") else "no" if sel in ("no", "en contra", "negativo") else "abstencion" if sel.startswith("abst") else "otros"
                    d[k] += 1
                    d["proyectos"].add(p["boletin"])
    fmt_votos = lambda d: [{"nombre": n, **{k: v for k, v in x.items() if k != "proyectos"}, "n_proyectos": len(x["proyectos"]),
                            "proyectos": sorted(x["proyectos"])} for n, x in sorted(d.items(), key=lambda kv: -(kv[1]["si"] + kv[1]["no"] + kv[1]["abstencion"]))]
    return {
        "autores_mociones": [{"nombre": n, "n_proyectos": len(b), "proyectos": sorted(b)} for n, b in sorted(autores.items(), key=lambda kv: (-len(kv[1]), kv[0]))],
        "autores_indicaciones_senado": [{"nombre": n, "n_indicaciones": d["n"], "n_proyectos": len(d["proyectos"]), "proyectos": sorted(d["proyectos"])}
                                         for n, d in sorted(indic.items(), key=lambda kv: -kv[1]["n"])],
        "votos_senadores": fmt_votos(votos_s),
        "votos_diputados": fmt_votos(votos_c),
        "nota": "Perfiles por fuente: los nombres se muestran como los informa cada fuente y no se cruzan entre sí.",
    }


# ----------------------------------------------------------------- Excel
def escribir_excel(ruta, proyectos, cambios, candidatos, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    azul = PatternFill("solid", fgColor="1F3864")
    hfont = Font(name="Verdana", bold=True, color="FFFFFF", size=9)
    bfont = Font(name="Verdana", size=9)

    def hoja(ws, campos, filas, anchos):
        ws.append(campos)
        for c in ws[1]:
            c.font, c.fill = hfont, azul
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for f in filas:
            ws.append([f.get(k, "") if not isinstance(f.get(k), bool) else ("Sí" if f.get(k) else "No") for k in campos])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = bfont
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(anchos):
            ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = w
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "Proyectos"
    hoja(ws, CAMPOS_PLANOS, proyectos, [11, 45, 10, 22, 50, 20, 25, 10, 22, 11, 14, 16, 28, 30, 26, 14, 26, 11, 8, 8,
                                        13, 11, 11, 12, 12, 18, 10, 10, 12, 18, 16, 12, 45, 12, 8, 9, 9, 9, 9, 9, 12, 9, 40, 9, 40, 6, 14, 14, 14])
    for nombre, archivo, anchos in [("Indicaciones", "indicaciones.csv", [11, 9, 12, 26, 8, 30, 13, 14, 40]),
                                    ("Indicaciones por autor", "indicaciones_por_autor.csv", [11, 40, 13, 12]),
                                    ("Hitos", "hitos.csv", [11, 12, 40, 12, 30]),
                                    ("Informes", "informes.csv", [11, 12, 60, 30, 40]),
                                    ("Votaciones Senado", "votaciones.csv", [11, 12, 9, 20, 28, 16, 6, 6, 9, 7, 11, 60]),
                                    ("Votaciones Cámara", "votaciones_camara.csv", [11, 9, 12, 7, 24, 18, 12, 6, 6, 9, 9, 60])]:
        filas = leer_csv(os.path.join(DATA, archivo))
        if filas:
            hoja(wb.create_sheet(nombre), list(filas[0]), filas, anchos)
    hoja(wb.create_sheet("Cambios"), ["fecha_deteccion", "boletin", "tipo", "antes", "despues"], cambios[-1000:][::-1], [17, 11, 22, 40, 60])
    hoja(wb.create_sheet("Candidatos"), ["boletin", "titulo", "fecha_deteccion", "motivo", "decision", "observacion"], candidatos, [11, 60, 12, 45, 12, 30])
    info = wb.create_sheet("Info")
    for k, v in meta.items():
        info.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v])
    wb.save(ruta)


# ----------------------------------------------------------------- principal
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sin-candidatos", action="store_true", help="no buscar proyectos nuevos")
    ap.add_argument("--sin-excel", action="store_true", help="no generar proyectos.xlsx")
    args = ap.parse_args()

    os.makedirs(DATA, exist_ok=True)
    os.makedirs(SALIDA, exist_ok=True)
    marca = ahora().strftime("%Y-%m-%d %H:%M")

    catalogo = {r["boletin"].strip(): r for r in leer_csv(os.path.join(CONFIG, "catalogo.csv")) if r.get("boletin", "").strip()}
    previos = {p["boletin"]: p for p in leer_json(os.path.join(DATA, "proyectos.json"), [])}
    print(f"[{marca}] Boletines en catálogo: {len(catalogo)}")

    def consultar(b):
        try:
            f = senado.ficha(b)
            return b, f, None if f else "El Senado no informa el boletín"
        except Exception as e:  # noqa: BLE001
            return b, None, f"{type(e).__name__}: {e}"

    with ThreadPoolExecutor(max_workers=6) as ex:
        resultados = list(ex.map(consultar, sorted(catalogo)))

    # boletines de indicaciones del Senado: se leen una sola vez y quedan en caché
    cache = cargar_cache_indicaciones()
    pendientes = {bi["iddocto"]: (b, bi) for b, f, _ in resultados if f for bi in f["boletines_indicaciones"]
                  if bi["iddocto"] and (cache.get(bi["iddocto"], {}).get("estado") != "leído"
                                        or any("texto" not in x for x in cache[bi["iddocto"]].get("items", [])))}
    if pendientes:
        with ThreadPoolExecutor(max_workers=3) as ex:
            for iddocto, res in zip(pendientes, ex.map(leer_boletin_indicaciones, [bi for _, bi in pendientes.values()])):
                cache[iddocto] = res
        print(f"  Boletines de indicaciones leídos ahora: {sum(cache[i]['estado'] == 'leído' for i in pendientes)} de {len(pendientes)}")
        guardar_cache_indicaciones(cache)

    nuevos, errores = {}, []
    for b, f, err in resultados:
        if f:
            nuevos[b] = registro(catalogo[b], f, cache)
        else:
            errores.append(f"{b}: {err}")
            if b in previos:                       # conserva el último dato bueno
                nuevos[b] = {**previos[b], **{k: catalogo[b].get(k, "") for k in CAMPOS_CATALOGO}}
    for e in errores:
        print("  ERROR", e)
    if catalogo and len(errores) / len(catalogo) > MAX_ERRORES:
        print(f"Demasiadas fallas ({len(errores)} de {len(catalogo)}). No se escribe nada.")
        sys.exit(1)

    # Cámara: votaciones de Sala (fuente independiente; si falla, se conserva lo ya descargado)
    fuentes_prev = leer_json(os.path.join(DATA, "meta.json"), {}).get("fuentes", {})
    estado_camara = dict(fuentes_prev.get("camara_votaciones", {}))
    camara_con_base = bool(estado_camara.get("ultima_consulta_exitosa"))
    try:
        cache_cam, err_det = actualizar_camara(sorted(nuevos))
        estado_camara.update({"estado": "ok" if not err_det else f"ok con {err_det} detalles pendientes",
                              "ultima_consulta_exitosa": marca[:10]})
        print(f"  Cámara: {len(cache_cam['votaciones'])} votaciones de proyectos seguidos ({err_det} detalles pendientes)")
    except Exception as e:  # noqa: BLE001
        cache_cam = leer_json(RUTA_CACHE_CAMARA, {"votaciones": {}, "detalles": {}})
        estado_camara["estado"] = f"sin conexión: {type(e).__name__}"[:80]
        print(f"  AVISO Cámara: no se pudo consultar ({type(e).__name__}: {e}). Se mantienen los datos anteriores.")
    hoy = marca[:10]
    for p in nuevos.values():
        agregar_camara(p, cache_cam)

    # agenda de la semana (citaciones, tabla de Sala, plazos de indicaciones) y agrupación ACOT
    agenda_prev = leer_json(RUTA_AGENDA, {})
    agenda, estado_agenda, citaciones_nuevas = agenda_mod.actualizar(agenda_prev, set(nuevos), hoy,
                                                                    sesiones_camara=lambda: camara.agenda_sala(hoy))
    print("  Agenda: " + ", ".join(f"{k} {v['estado']}" for k, v in estado_agenda.items()) + f" · {len(agenda['items'])} sesiones")
    acot = cargar_acot()
    for p in nuevos.values():
        p["agenda"], p["plazo_indicaciones"] = agenda_mod.agenda_de(p["boletin"], agenda)
        p.update(acot.get(p["boletin"], {"acot_n": "", "acot_iniciativa": "", "acot_grupo": ""}))
        gestion.evaluar(p, hoy)

    orden = sorted(nuevos.values(), key=lambda p: int(p["boletin"].split("-")[0]))

    # cambios
    cambios_nuevos = detectar_cambios(previos, nuevos, marca, camara_con_base) if previos else []
    if agenda_prev.get("items"):                    # la primera carga de la agenda no genera avisos
        for x in citaciones_nuevas:
            for b in x["seguidos"]:
                cambios_nuevos.append({"fecha_deteccion": marca, "boletin": b, "tipo": "Citado a sesión",
                                       "antes": "", "despues": f"{x['fecha']} {x.get('horario', '')} – {x['instancia']}".strip()})
    hist = leer_csv(os.path.join(DATA, "cambios.csv")) + cambios_nuevos
    print(f"  Cambios detectados: {len(cambios_nuevos)}")

    # candidatos
    candidatos = leer_csv(os.path.join(DATA, "candidatos.csv"))
    alerta = os.path.join(SALIDA, "alerta_candidatos.md")
    if os.path.exists(alerta):
        os.remove(alerta)
    if not args.sin_candidatos:
        try:
            revisados = {c["boletin"] for c in candidatos}
            nuevos_c = buscar_candidatos(set(nuevos), revisados, marca)
            candidatos += nuevos_c
            print(f"  Candidatos nuevos: {len(nuevos_c)}")
            if nuevos_c:
                with open(alerta, "w", encoding="utf-8") as f:
                    f.write("Proyectos de seguridad con acción del actual Gobierno que no están en `config/catalogo.csv`.\n\n")
                    f.write("Para incorporarlos, agregar una fila en `config/catalogo.csv`. Para descartarlos, cambiar `decision` en `data/candidatos.csv`.\n\n")
                    f.write("| Boletín | Título | Motivo |\n|---|---|---|\n")
                    for c in nuevos_c:
                        url = f"https://tramitacion.senado.cl/appsenado/templates/tramitacion/index.php?boletin_ini={c['boletin']}"
                        f.write(f"| [{c['boletin']}]({url}) | {c['titulo']} | {c['motivo']} |\n")
        except Exception as e:  # noqa: BLE001
            print(f"  AVISO: no se pudo revisar candidatos ({type(e).__name__}: {e})")

    # escritura (solo cambia el contenido si cambiaron los datos)
    antes = {n: contenido(os.path.join(DATA, n)) for n in ("proyectos.json", "cambios.csv", "candidatos.csv", "indicaciones.csv", "boletines_indicaciones.csv", "votaciones_camara.csv")}
    escribir_json(os.path.join(DATA, "proyectos.json"), orden)
    escribir_json(os.path.join(DATA, "tablero.json"), [{k: p.get(k, "") for k in CAMPOS_TABLERO} for p in orden])
    escribir_csv(os.path.join(DATA, "proyectos.csv"), orden, CAMPOS_PLANOS)
    escribir_csv(os.path.join(DATA, "tramites.csv"),
                 [{"boletin": p["boletin"], **t} for p in orden for t in p["tramites"]],
                 ["boletin", "fecha", "sesion", "camara", "etapa", "descripcion"])
    escribir_csv(os.path.join(DATA, "urgencias.csv"),
                 [{"boletin": p["boletin"], **u} for p in orden for u in p["urgencias"]],
                 ["boletin", "tipo", "ingreso", "retiro", "camara", "mensaje"])
    escribir_detalle(orden, cache)
    vc = [{"boletin": p["boletin"], **{k: v[k] for k in v if k != "votos"}} for p in orden for v in p.get("votaciones_camara", [])]
    escribir_csv(os.path.join(DATA, "votaciones_camara.csv"), vc,
                 ["boletin", "id", "fecha", "hora", "tipo", "quorum", "resultado", "si", "no", "abstencion", "dispensados", "descripcion"])
    escribir_csv(os.path.join(DATA, "votos_camara.csv"),
                 [{"boletin": p["boletin"], "votacion_id": v["id"], "fecha": v["fecha"], **x} for p in orden for v in p.get("votaciones_camara", []) for x in v["votos"]],
                 ["boletin", "votacion_id", "fecha", "diputado_id", "diputado", "seleccion"])
    escribir_json(os.path.join(DATA, "actores.json"), calcular_actores(orden))
    escribir_json(RUTA_AGENDA, agenda)
    escribir_json(os.path.join(DATA, "acot.json"), acot_lista(acot_filas(), nuevos))
    escribir_json(os.path.join(SALIDA, "cambios_nuevos.json"), cambios_nuevos)
    escribir_csv(os.path.join(DATA, "cambios.csv"), hist, ["fecha_deteccion", "boletin", "tipo", "antes", "despues"])
    escribir_json(os.path.join(DATA, "cambios.json"), hist[-500:][::-1])
    escribir_csv(os.path.join(DATA, "candidatos.csv"), candidatos, ["boletin", "titulo", "fecha_deteccion", "motivo", "decision", "observacion"])
    hubo_cambio = any(contenido(os.path.join(DATA, n)) != antes[n] for n in antes)

    meta_prev = leer_json(os.path.join(DATA, "meta.json"), {})
    meta = {
        "fuente": "Servicio web de tramitación del Senado (tramitacion.senado.cl/wspublico)",
        "fecha_verificacion": marca[:10],
        "ultima_actualizacion_datos": marca if hubo_cambio or not meta_prev else meta_prev.get("ultima_actualizacion_datos", marca),
        "n_proyectos": len(orden),
        "n_en_tramitacion": sum(p["situacion"] == "En tramitación" for p in orden),
        "inicio_gobierno": INICIO_GOBIERNO,
        "fuentes": {
            "senado": {"estado": "ok" if not errores else f"ok con {len(errores)} boletines sin respuesta", "ultima_consulta_exitosa": marca[:10]},
            "camara_votaciones": estado_camara,
            "agenda": {"estado": "ok" if all(v["estado"] == "ok" for v in estado_agenda.values()) else
                       ("parcial" if any(v["estado"] == "ok" for v in estado_agenda.values()) else "sin conexión"),
                       "detalle": estado_agenda,
                       "ultima_consulta_exitosa": marca[:10] if any(v["estado"] == "ok" for v in estado_agenda.values())
                       else meta_prev.get("fuentes", {}).get("agenda", {}).get("ultima_consulta_exitosa", "")},
        },
    }
    escribir_json(os.path.join(DATA, "meta.json"), meta)

    if not args.sin_excel and (hubo_cambio or not os.path.exists(os.path.join(DATA, "proyectos.xlsx"))):
        escribir_excel(os.path.join(DATA, "proyectos.xlsx"), orden, hist, candidatos, meta)
        print("  Excel regenerado")
    print(f"  Listo. Proyectos: {len(orden)} | errores: {len(errores)} | datos cambiaron: {hubo_cambio}")


if __name__ == "__main__":
    main()
